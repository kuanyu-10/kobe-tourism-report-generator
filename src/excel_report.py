from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Border,
    Side,
    Alignment
)

def adjust_column_width(ws):

    for column_cells in ws.columns:

        max_length = 0

        column_letter = column_cells[0].column_letter

        for cell in column_cells:

            if cell.value is not None:

                value_length = len(str(cell.value))

                if value_length > max_length:
                    max_length = value_length

        ws.column_dimensions[column_letter].width = max_length + 2

def apply_border(ws):

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    for row in ws.iter_rows():

        for cell in row:

            if cell.value is not None:

                cell.border = thin_border


def apply_alignment(ws):

    center_alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    for row in ws.iter_rows():

        for cell in row:

            if cell.value is not None:

                cell.alignment = center_alignment


def create_excel_report(
    kpi,
    area_df,
    category_df,
    top10_df
):

    wb = Workbook()

    ws = wb.active
    ws.title = "KPI Summary"

    ws["A1"] = "KPI"
    ws["B1"] = "Value"

    header_fill = PatternFill(
    fill_type="solid",
    fgColor="4472C4"
)

    ws["A1"].font = Font(bold=True)
    ws["B1"].font = Font(bold=True)
    ws["A1"].fill = header_fill
    ws["B1"].fill = header_fill


    ws["A2"] = "総スポット数"
    ws["B2"] = kpi["total_spots"]

    ws["A3"] = "平均評価"
    ws["B3"] = kpi["average_rating"]

    ws["A4"] = "総レビュー数"
    ws["B4"] = kpi["total_reviews"]

    ws["A5"] = "人気スポット"
    ws["B5"] = kpi["top_spot"]

    # Area Analysis Sheet
    area_ws = wb.create_sheet("Area Analysis")

    area_ws.append([
        "エリア",
        "スポット数",
        "平均評価",
        "総レビュー数"
    ])
    for cell in area_ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for _, row in area_df.iterrows():
        area_ws.append([
            row["area"],
            row["spot_count"],
            row["avg_rating"],
            row["total_reviews"]
        ])

    category_ws = wb.create_sheet("Category Analysis")

    category_ws.append([
        "カテゴリ",
        "スポット数",
        "平均評価",
        "総レビュー数",
        "平均Popularity Score"
    ])
    for cell in category_ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for _, row in category_df.iterrows():

        category_ws.append([
            row["category"],
            row["spot_count"],
            row["avg_rating"],
            row["total_reviews"],
            row["avg_popularity_score"]
        ])


    # Top10 Ranking Sheet
    top10_ws = wb.create_sheet("Top10 Ranking")

    top10_ws.append([
        "順位",
        "スポット名",
        "エリア",
        "カテゴリ",
        "評価",
        "レビュー数",
        "Popularity Score"
    ])
    for cell in top10_ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for rank, (_, row) in enumerate(
            top10_df.iterrows(),
            start=1
            ):

        top10_ws.append([
            rank,
            row["spot_name"],
            row["area"],
            row["category"],
            row["rating"],
            row["review_count"],
            row["popularity_score"]
        ])

    for ws in wb.worksheets:
        adjust_column_width(ws)
        apply_border(ws)
        apply_alignment(ws)


    wb.save("output/kobe_tourism_report.xlsx")